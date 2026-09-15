"""Gera o guia didático em Word, sem modificar o código da aplicação."""
from pathlib import Path
from textwrap import dedent
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Guia_CRUD_Categorias_Laravel.docx'
doc = Document()

# Preset: compact_reference_guide. Cabeçalho: editorial_cover compacto.
# Exceções nomeadas: Code (Consolas 9,5 pt), File (Consolas 9 pt),
# Title (24 pt) e Subtitle (12 pt). Código sem cores para facilitar a cópia.
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

def style(name, size, before, after, color='202020', font='Calibri', line=1.25, bold=False):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name, s.font.size, s.font.bold = font, Pt(size), bold
    s.font.color.rgb = RGBColor.from_string(color)
    f = s.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    f.widow_control = True
    return s

style('Normal', 11, 0, 6)
style('Title', 24, 0, 8, '1F4D78', bold=True)
style('Subtitle', 12, 0, 12, '555555')
for name, size, before, after, color in [
    ('Heading 1', 16, 18, 10, '2E74B5'),
    ('Heading 2', 13, 14, 7, '2E74B5'),
    ('Heading 3', 12, 10, 5, '1F4D78'),
]:
    style(name, size, before, after, color, bold=True).paragraph_format.keep_with_next = True
style('Code', 9.5, 0, 0, font='Consolas', line=1.0)
style('File', 9, 0, 7, '555555', font='Consolas', line=1.0).paragraph_format.keep_with_next = True
style('Note', 11, 6, 8, '1F4D78')
style('Header', 9, 0, 0, '777777', line=1.0)
style('Footer', 9, 0, 0, '777777', line=1.0)
style('List Bullet', 11, 0, 4)

# Uma lista real do Word, com recuo definido pelo preset.
numbering = doc.part.numbering_part.element
abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), '42')
level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
for tag, value in [('start', '1'), ('numFmt', 'bullet'), ('lvlText', '•'), ('lvlJc', 'left')]:
    e = OxmlElement('w:' + tag); e.set(qn('w:val'), value); level.append(e)
ppr = OxmlElement('w:pPr')
tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), '540'); tabs.append(tab); ppr.append(tabs)
ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '540'); ind.set(qn('w:hanging'), '271'); ppr.append(ind)
level.append(ppr); abstract.append(level); numbering.append(abstract)
num = OxmlElement('w:num'); num.set(qn('w:numId'), '42')
aid = OxmlElement('w:abstractNumId'); aid.set(qn('w:val'), '42'); num.append(aid); numbering.append(num)

sec.header.paragraphs[0].text = 'LARAVEL • CRUD DE CATEGORIAS'
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Página ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'CRUD de categorias com Laravel: passo a passo para iniciantes'
doc.core_properties.subject = 'Guia do CRUD implementado no projeto, seguindo MVC'
doc.core_properties.author = ''
lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'pt-BR')
doc.styles['Normal'].element.get_or_add_rPr().append(lang)

def p(text, sty=None):
    return doc.add_paragraph(text, sty)

def bullet(text):
    para = p(text, 'List Bullet')
    np = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
    ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '42')
    np.extend([ilvl, ni]); para._p.get_or_add_pPr().append(np)

def code(text):
    lines = dedent(text).strip('\n').splitlines()
    for i, line in enumerate(lines):
        para = p(line, 'Code')
        para.paragraph_format.keep_with_next = i < len(lines) - 1
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F4F6F9')
        para._p.get_or_add_pPr().append(shade)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def h(text):
    doc.add_heading(text, 2)

def file(text):
    p(text, 'File')

p('CRUD de categorias\ncom Laravel', 'Title')
p('Passo a passo para iniciantes', 'Subtitle')
p('Neste guia, você vai entender como cadastrar, listar, visualizar, editar e excluir categorias no projeto que construímos.')
h('O que é CRUD?')
bullet('Create: criar uma categoria. Exemplo: cadastrar “Livros”.')
bullet('Read: consultar as categorias na lista ou na tela de detalhes.')
bullet('Update: atualizar uma categoria. Exemplo: mudar o nome.')
bullet('Delete: excluir uma categoria pelo botão da listagem.')
h('O que é MVC?')
bullet('Model: conversa com o banco de dados. Aqui, é Categoria.php.')
bullet('View: mostra a tela para o usuário. Aqui, são os arquivos Blade.')
bullet('Controller: recebe o pedido do usuário e decide o que fazer.')
p('Exemplo: ao clicar em Salvar, a rota chama o Controller. Ele valida os campos, usa o Model para salvar e volta para a listagem.')
h('Como acompanhar')
p('Abra o projeto no editor e siga os arquivos indicados. Este guia parte do Laravel já configurado, com banco, login e os componentes Flux do projeto.')
p('Os arquivos já existem: não é necessário recriá-los. Os trechos mostram as partes principais; algumas classes de aparência foram omitidas para facilitar a leitura. Métodos do Controller devem ficar dentro da classe CategoriaController.', 'Note')

page('Passo 1 • Criar a tabela em português')
p('Uma migration é um arquivo que descreve a estrutura de uma tabela. Nossa tabela se chama categorias.')
file('database/migrations/2026_09_15_000000_criar_tabela_categorias.php')
p('Dentro de up(), criamos as colunas:')
code('''
public function up(): void
{
    Schema::create('categorias', function (Blueprint $table) {
        $table->id();
        $table->string('nome', 100);
        $table->text('descricao')->nullable();
        $table->timestamp('criado_em')->nullable();
        $table->timestamp('atualizado_em')->nullable();
    });
}
''')
bullet('id: número que identifica cada categoria, gerado pelo banco.')
bullet('nome: texto obrigatório, com até 100 caracteres.')
bullet('descricao: texto opcional. nullable() permite deixar sem valor.')
bullet('criado_em e atualizado_em: datas controladas pelo Laravel.')
p('O método down() desfaz esta migration, removendo a tabela:')
code('''
public function down(): void
{
    Schema::dropIfExists('categorias');
}
''')
h('Aplicar a migration')
p('No terminal, dentro da pasta do projeto:')
code('php artisan migrate')
p('Esse comando executa as migrations pendentes. A nossa já foi aplicada no projeto.')
p('Se estiver repetindo o exercício em outro projeto preparado, crie o arquivo com php artisan make:migration criar_tabela_categorias e edite os métodos acima. O prefixo de data será diferente.', 'Note')

page('Passo 2 • Preparar o Model e as rotas')
file('app/Models/Categoria.php')
p('O Model representa a categoria no banco. O Eloquent é a ferramenta do Laravel que permite trabalhar com os registros usando PHP.')
code('''
<?php

namespace App\\Models;

use Illuminate\\Database\\Eloquent\\Model;

class Categoria extends Model
{
    protected $table = 'categorias';
    protected $fillable = ['nome', 'descricao'];

    public const CREATED_AT = 'criado_em';
    public const UPDATED_AT = 'atualizado_em';
}
''')
p('$table define a tabela. $fillable indica quais campos create() e update() podem preencher. As constantes dizem ao Laravel os nomes das colunas de data.')
file('routes/web.php')
p('Uma rota liga um endereço ao método do Controller. No início do arquivo, importamos o Controller:')
code('use App\\Http\\Controllers\\CategoriaController;')
p('Dentro do grupo de rotas que já usa auth e verified, adicionamos:')
code('''
Route::resource('categorias', CategoriaController::class)
    ->parameters(['categorias' => 'categoria']);
''')
p('resource cria as sete rotas do CRUD. O parâmetro categoria corresponde a Categoria $categoria nos métodos. auth exige login; verified aplica a verificação de e-mail quando ela estiver habilitada para o usuário.')
p('Em outro projeto preparado, o comando php artisan make:model Categoria cria o Model inicial. No nosso projeto, ele já está pronto.', 'Note')

page('Passo 3 • Abrir as telas pelo Controller')
file('app/Http/Controllers/CategoriaController.php')
p('O início do arquivo identifica a pasta lógica da classe (namespace) e importa as classes que serão usadas (use):')
code('''
<?php

namespace App\\Http\\Controllers;

use App\\Models\\Categoria;
use Illuminate\\Http\\Request;

class CategoriaController extends Controller
{
    // Os métodos das etapas 3, 4 e 5 ficam aqui.
}
''')
p('Estes quatro métodos abrem as telas:')
code('''
public function index()
{
    $categorias = Categoria::orderBy('id')->paginate(10);
    return view('categorias.index', compact('categorias'));
}

public function create()
{
    return view('categorias.create');
}

public function show(Categoria $categoria)
{
    return view('categorias.show', compact('categoria'));
}

public function edit(Categoria $categoria)
{
    return view('categorias.edit', compact('categoria'));
}
''')
p('index busca os registros por código, com 10 por página. compact envia a variável para a tela. view abre o arquivo Blade: categorias.index significa categorias/index.blade.php.')
p('Em show e edit, o Laravel encontra a categoria pelo ID da URL. Se ela não existir, retorna 404. Em outro projeto, php artisan make:controller CategoriaController --resource gera os métodos iniciais.')

page('Passo 4 • Validar e cadastrar')
file('app/Http/Controllers/CategoriaController.php')
p('Request contém os dados enviados pelo formulário. Antes de salvar, verificamos se os valores são válidos. Este método será usado no cadastro e na edição:')
code('''
private function validar(Request $request)
{
    return $request->validate([
        'nome' => ['required', 'string', 'max:100'],
        'descricao' => ['nullable', 'string', 'max:1000'],
    ], [
        'nome.required' => 'Informe o nome da categoria.',
        'nome.string' => 'O nome deve ser um texto.',
        'nome.max' =>
            'O nome deve ter no máximo 100 caracteres.',
        'descricao.string' => 'A descrição deve ser um texto.',
        'descricao.max' =>
            'A descrição deve ter no máximo 1000 caracteres.',
    ]);
}
''')
bullet('required: o campo precisa ser preenchido.')
bullet('string: o valor precisa ser um texto.')
bullet('max: limita a quantidade de caracteres do texto.')
bullet('nullable: aceita que a descrição fique sem valor.')
p('Se a validação falhar, o Laravel volta ao formulário com os erros e os valores digitados. Se passar, store salva a categoria:')
code('''
public function store(Request $request)
{
    Categoria::create($this->validar($request));

    return redirect()->route('categorias.index')
        ->with('sucesso', 'Categoria cadastrada com sucesso!');
}
''')
p('$this->validar(...) chama a validação desta classe. create(...) grava uma nova linha. redirect volta para a listagem. with guarda a mensagem de sucesso para a próxima requisição.')

page('Passo 5 • Atualizar e excluir')
file('app/Http/Controllers/CategoriaController.php')
h('Atualizar uma categoria')
p('update recebe os dados do formulário e a categoria que será alterada. Usamos a mesma validação do cadastro:')
code('''
public function update(Request $request, Categoria $categoria)
{
    $categoria->update($this->validar($request));

    return redirect()->route('categorias.index')
        ->with('sucesso', 'Categoria atualizada com sucesso!');
}
''')
p('Exemplo: a categoria de código 1 se chama “Livros”. Ao salvar “Material escolar”, update altera essa mesma linha do banco.')
h('Excluir uma categoria')
p('destroy recebe a categoria encontrada pelo Laravel e remove o registro:')
code('''
public function destroy(Categoria $categoria)
{
    $categoria->delete();

    return redirect()->route('categorias.index')
        ->with('sucesso', 'Categoria excluída com sucesso!');
}
''')
p('O botão Excluir fica no index.blade.php. Ele envia um formulário para destroy. Não precisamos de uma tela nem de um arquivo excluir.blade.php.')
h('Uma diferença importante')
bullet('create abre o formulário; store salva uma nova categoria.')
bullet('edit abre o formulário preenchido; update salva a alteração.')
bullet('show mostra os detalhes; destroy exclui o registro.')
p('Os nomes dos métodos seguem a convenção do Laravel. O nome da tabela, os campos e as mensagens do nosso CRUD estão em português.')

page('Passo 6 • Reutilizar o layout e o formulário')
file('resources/views/layouts/app.blade.php')
p('O layout já existe no projeto. Ele reúne a estrutura visual da aplicação. Nas telas, usamos <x-layouts::app>; o conteúdo é colocado no {{ $slot }} do layout.')
code('''
<x-layouts::app title="Categorias">
    <h1>Categorias</h1>
    <!-- O conteúdo desta tela fica aqui. -->
</x-layouts::app>
''')
p('Não usamos @extends neste caso, porque o layout foi preparado como componente Blade. Também não precisamos de um layout separado na pasta categorias.')
file('resources/views/categorias/form.blade.php')
p('O cadastro e a edição compartilham os mesmos campos. Trecho sem as classes de aparência:')
code('''
@csrf

<flux:input name="nome" label="Nome"
    :value="old('nome', $categoria->nome ?? '')"
    required maxlength="100" />
@error('nome')
    <p>{{ $message }}</p>
@enderror

<flux:textarea name="descricao" label="Descrição (opcional)"
    rows="4" maxlength="1000"
>{{ old('descricao', $categoria->descricao ?? '') }}</flux:textarea>
@error('descricao')
    <p>{{ $message }}</p>
@enderror

<flux:button type="submit">Salvar</flux:button>
<a href="{{ route('categorias.index') }}">Cancelar</a>
''')
p('@csrf adiciona um token de proteção ao formulário. old recupera o que foi digitado após um erro. Na edição, usamos os dados da categoria; ?? fornece um valor alternativo quando não há dado. @error mostra a mensagem de validação.')
p('flux:input, flux:textarea e flux:button são componentes visuais já instalados neste projeto. Os campos precisam dos atributos name para enviar os valores corretos.')

page('Passo 7 • Criar as telas de cadastro e edição')
file('resources/views/categorias/create.blade.php')
p('O formulário de cadastro envia os dados por POST para a rota categorias.store:')
code('''
<x-layouts::app title="Nova categoria">
    <h1>Nova categoria</h1>

    <form action="{{ route('categorias.store') }}" method="POST">
        @include('categorias.form')
    </form>
</x-layouts::app>
''')
file('resources/views/categorias/edit.blade.php')
p('Na edição, enviamos também a categoria na rota. O Laravel usa o ID dela para montar o endereço:')
code('''
<x-layouts::app title="Editar categoria">
    <h1>Editar categoria</h1>

    <form action="{{ route('categorias.update', $categoria) }}"
        method="POST">
        @method('PUT')
        @include('categorias.form')
    </form>
</x-layouts::app>
''')
bullet('@include insere o conteúdo de form.blade.php nas duas telas.')
bullet('O @csrf já está dentro do formulário compartilhado.')
bullet('@method(\'PUT\') informa ao Laravel que queremos atualizar.')
p('Formulários HTML enviam GET ou POST. A diretiva @method cria um campo escondido para o Laravel interpretar o envio como PUT. Na exclusão, fazemos o mesmo com DELETE.')
h('Como os dados chegam ao Controller')
p('Ao clicar em Salvar, o navegador envia nome e descricao. A rota escolhe store ou update, que recebe esses dados por Request e faz a validação antes de gravar.')

page('Passo 8 • Mostrar os detalhes e a listagem')
file('resources/views/categorias/show.blade.php')
p('show recebe uma categoria e mostra seus dados. Versão com HTML simplificado:')
code('''
<x-layouts::app title="Detalhes da categoria">
    <h1>Detalhes da categoria</h1>
    <p>Código: {{ $categoria->id }}</p>
    <p>Nome: {{ $categoria->nome }}</p>
    <p>{{ $categoria->descricao ?? 'Sem descrição' }}</p>
    <a href="{{ route('categorias.edit', $categoria) }}">Editar</a>
    <a href="{{ route('categorias.index') }}">Voltar</a>
</x-layouts::app>
''')
file('resources/views/categorias/index.blade.php')
p('A listagem usa o mesmo layout. Dentro dele, mostramos a mensagem de sucesso, o botão Nova categoria e a tabela:')
code('''
@if (session('sucesso'))
    <p>{{ session('sucesso') }}</p>
@endif

<a href="{{ route('categorias.create') }}">Nova categoria</a>

<table>
    <thead>
        <tr>
            <th>Código</th><th>Nome</th>
            <th>Descrição</th><th>Ações</th>
        </tr>
    </thead>
    <tbody>
        <!-- Insira aqui o trecho da próxima etapa. -->
    </tbody>
</table>
''')
p('session(\'sucesso\') lê a mensagem enviada pelo Controller. As chaves {{ }} mostram valores como texto seguro, sem executar HTML que o usuário tenha digitado.')
p('O próximo trecho fica dentro de tbody e cria uma linha para cada categoria. Os exemplos desta etapa mostram a lógica; as classes do projeto cuidam da aparência.')

page('Passo 9 • Colocar as ações no index')
file('resources/views/categorias/index.blade.php — dentro de tbody')
p('@forelse percorre as categorias. Dentro da coluna Ações, ficam os links Ver e Editar e o formulário Excluir:')
code('''
@forelse ($categorias as $categoria)
    <tr>
        <td>{{ $categoria->id }}</td>
        <td>{{ $categoria->nome }}</td>
        <td>{{ $categoria->descricao ?? 'Sem descrição' }}</td>
        <td>
            <a href="{{ route('categorias.show', $categoria) }}">
                Ver
            </a>
            <a href="{{ route('categorias.edit', $categoria) }}">
                Editar
            </a>

            <form
                action="{{ route('categorias.destroy', $categoria) }}"
                method="POST"
                onsubmit="return confirm('Deseja excluir esta categoria?')">
                @csrf
                @method('DELETE')
                <button type="submit">Excluir</button>
            </form>
        </td>
    </tr>
@empty
    <tr>
        <td colspan="4">Nenhuma categoria cadastrada.</td>
    </tr>
@endforelse
''')
bullet('@empty define o que aparece quando não há categorias.')
bullet('@method(\'DELETE\') faz a rota chamar destroy.')
bullet('confirm abre a confirmação. Cancelar impede o envio.')
bullet('@csrf protege o formulário de exclusão, assim como o de cadastro.')
p('O formulário fica diretamente no index, dentro de cada linha. Assim, cada botão Excluir aponta para a categoria daquela linha. A exclusão remove o registro do banco.')

page('Passo 10 • Paginar, acessar e testar')
h('Paginação e menu')
p('O Controller usa paginate(10). Abaixo da tabela, o index exibe Anterior e Próxima. Exemplo dos links usados, mantendo o texto em português:')
code('''
@if ($categorias->previousPageUrl())
    <a href="{{ $categorias->previousPageUrl() }}">Anterior</a>
@endif
@if ($categorias->nextPageUrl())
    <a href="{{ $categorias->nextPageUrl() }}">Próxima</a>
@endif
''')
p('No projeto, esse trecho fica em uma navegação que também mostra a página atual. O menu Categorias foi adicionado em resources/views/layouts/app/sidebar.blade.php, apontando para route(\'categorias.index\').')
h('Abrir o projeto')
p('Com as dependências já instaladas, execute os comandos na pasta do projeto. O primeiro prepara os estilos; o segundo inicia o servidor:')
code('''
npm run build
php artisan serve
''')
p('Abra http://localhost:8000, faça login e clique em Categorias. Também é possível acessar http://localhost:8000/categorias após o login. Se o terminal indicar outra porta, use a porta informada.')
h('Conferir as quatro operações')
bullet('Cadastrar: clique em Nova categoria, digite “Livros” e salve.')
bullet('Consultar: veja “Livros” na lista e clique em Ver.')
bullet('Editar: mude o nome para “Material escolar” e salve.')
bullet('Validar: tente salvar sem nome. Deve aparecer um erro.')
bullet('Excluir: clique em Excluir. Teste Cancelar e depois confirme.')
p('Para rodar os testes automatizados do CRUD:')
code('php artisan test --filter=CategoriaTest')
p('Resumo: a rota recebe o pedido; o Controller organiza a ação; o Model trabalha com o banco; a View mostra o resultado. Esse é o fluxo MVC do nosso CRUD.', 'Note')

ROOT.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

# Verificação estrutural: reabre o DOCX e confirma seções e arquivos citados.
loaded = Document(OUT)
text = '\n'.join(p.text for p in loaded.paragraphs)
assert sum(p.style.name == 'Heading 1' for p in loaded.paragraphs) == 10
for token in ['criado_em', 'atualizado_em', "@method('DELETE')", '<x-layouts::app', 'CategoriaController']:
    assert token in text, token
assert loaded.sections[0].page_width.twips == 12240
assert loaded.sections[0].left_margin.twips == 1440
with ZipFile(OUT) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):
            etree.fromstring(z.read(name))
print(f'Documento criado e estrutura verificada: {OUT}')
print(f'Parágrafos: {len(loaded.paragraphs)}; etapas: 10')
